#!/usr/bin/env python3
"""Actualise la présentation du groupe sans changer son identité visuelle.

Dépendance : python-pptx. Exécuter depuis n’importe quel dossier.
Les résultats affichés correspondent à docs/RESULTATS_VALIDATION.md.
"""
from copy import deepcopy
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData


def replace_text(shape, text):
    """Conserve les propriétés du premier paragraphe et de son texte."""
    frame = shape.text_frame
    paragraph = frame.paragraphs[0]
    paragraph_style = deepcopy(paragraph._p.pPr)
    run_style = deepcopy(paragraph.runs[0]._r.rPr) if paragraph.runs else None
    frame.clear()
    for index, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        if paragraph_style is not None:
            if paragraph._p.pPr is not None:
                paragraph._p.remove(paragraph._p.pPr)
            paragraph._p.insert(0, deepcopy(paragraph_style))
        run = paragraph.add_run()
        run.text = line
        if run_style is not None:
            run._r.insert(0, deepcopy(run_style))


def main():
    path = Path(__file__).resolve().parents[1] / "presentation/PRESENTATION_SOAP_CLIENTS.pptx"
    deck = Presentation(path)
    changes = {
        (3, 26): "Santé, métriques, WSDL et consultation SOAP périodique",
        (3, 27): "11 services Compose · secrets générés localement · ports locaux",
        (5, 19): "Sur une base fraîche : relancer l’import ne crée aucun doublon.",
        (7, 19): "WSDL 1.1 · document/literal · SOAP 1.1\n23 tests SOAP sur pile Docker réelle · requêtes SoapUI fournies",
        (8, 2): "Fault : HTTP 500. Détail métier typé par le WSDL ; une erreur d’enveloppe peut ne pas avoir de détail métier.",
        (9, 10): "ignore les fiches SOAP\nmême après réconciliation",
        (10, 25): "14 éléments, 8 déclencheurs, 2 graphiques. WSDL, appel SOAP et panne MongoDB testés sur Zabbix 7.0.32.",
        (11, 3): "23 / 23",
        (11, 4): "tests SOAP sur pile Docker réelle",
        (11, 6): "33 / 33",
        (11, 7): "tests Python ETL / réconciliation",
        (11, 9): "27 / 27",
        (11, 10): "contrôles de la chaîne complète",
        (11, 11): "Zabbix : 9 contrôles, panne MongoDB détectée puis résolue. Couverture Python : 98 %.",
        (12, 0): "Groupe 5 SI : répartition du travail",
        (12, 4): "KOANDA",
        (12, 5): "GILDAS — Infrastructure Docker, dépôt Git et README",
        (12, 9): "SANKARA",
        (12, 10): "Michail Alioun Juste — MySQL, jeu de données et DAG d’import",
        (12, 14): "BARRY",
        (12, 15): "EL NOUR — Service SOAP Laravel, WSDL et validation",
        (12, 19): "CONGO",
        (12, 20): "Bénéré Nelson Cephas — MongoDB, DAG de réconciliation",
        (12, 24): "KONE",
        (12, 25): "Assita — Zabbix, tests SoapUI et documentation",
        (12, 26): "IBAM · Groupe 5 SI · Exercice 7 · 2026",
        (13, 2): "Limites : délai lié aux DAG et aux reprises · accès local sans authentification SOAP · pas de création/suppression",
    }
    for (slide_number, shape_index), text in changes.items():
        replace_text(deck.slides[slide_number - 1].shapes[shape_index], text)
    for shape in deck.slides[4].shapes:
        if shape.has_table:
            replace_text(shape.table.cell(2, 1), "mise à jour si la fiche reste gérée par MySQL")
            replace_text(shape.table.cell(4, 1), "jamais écrasé, même après retour vers MySQL")
    for shape in deck.slides[10].shapes:
        if shape.has_chart:
            data = CategoryChartData()
            data.categories = ["WSDL", "getClientDetails", "searchClients", "updateClient", "Enveloppe", "Panne base"]
            data.add_series("Nombre de cas", (1, 3, 7, 8, 2, 2))
            shape.chart.replace_data(data)
            replace_text(shape.chart.chart_title, "23 cas de test SOAP par domaine")
    deck.core_properties.title = "Service SOAP de clients — Groupe 5 SI"
    deck.core_properties.author = "SANKARA, KOANDA, BARRY, CONGO, KONE — Groupe 5 SI"
    deck.core_properties.subject = "Exercice 7 — Laravel, Airflow, MySQL, MongoDB et Zabbix"
    deck.save(path)
    print(f"Présentation actualisée : {len(deck.slides)} diapositives.")


if __name__ == "__main__":
    main()
